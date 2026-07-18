"""Render report instance to DOCX (docxtpl + tenant template path)."""

from __future__ import annotations

import json
import re
import shutil
import tempfile
import unicodedata
from pathlib import Path
from typing import Any

from duckclaw.report_engine.state import (
    assign_context_value,
    build_render_context,
    image_render_specs,
)


def _lenient_jinja_env() -> Any:
    """Jinja env que no crashea con huecos faltantes/anidados.

    La plantilla suele tener `{{ grupo.item }}`; si una sección no se rellenó,
    el `Undefined` por defecto revienta en el acceso a la clave y tumba TODO el
    render (el usuario ve un fallo y su agente cae a pandoc, perdiendo formato).
    `ChainableUndefined` deja `{{ a.b.c }}` → '' sin error.
    """
    import jinja2

    undefined = getattr(jinja2, "ChainableUndefined", jinja2.Undefined)
    return jinja2.Environment(undefined=undefined, autoescape=False)


def assert_template_uri_readable(template_uri: str, allowed_roots: list[Path]) -> Path:
    """Re-valida que la plantilla sigue bajo raíces permitidas (ALLOWED ∪ OUTPUT)."""
    template_path = Path(template_uri).expanduser().resolve()
    if not template_path.is_file():
        raise ValueError(f"Plantilla no accesible: {template_uri}")
    if not allowed_roots:
        raise ValueError("No hay raíces de conocimiento configuradas para leer la plantilla")
    ok = any(
        template_path == root.resolve() or root.resolve() in template_path.parents
        for root in allowed_roots
    )
    if not ok:
        raise ValueError(
            "template_uri fuera de DUCKCLAW_KNOWLEDGE_ALLOWED_ROOTS / OUTPUT_ROOTS. "
            "Vuelve a registrar la plantilla desde una ruta permitida."
        )
    return template_path


# Página carta ~11" con márgenes ~1": alto útil ~9". Google Docs a menudo
# oculta InlineImage si cy supera el alto de página (deja el hueco en blanco).
_MAX_IMAGE_WIDTH_IN = 6.0
_MAX_IMAGE_HEIGHT_IN = 7.0


def _resolve_image_path(raw_path: str, image_roots: list[Path]) -> Path:
    """Valida que la imagen exista y viva bajo una raíz permitida (anti path-traversal)."""
    candidate = Path(raw_path).expanduser().resolve()
    if not candidate.is_file():
        raise ValueError(f"Imagen no accesible: {raw_path}")
    if image_roots:
        ok = any(
            candidate == root.resolve() or root.resolve() in candidate.parents
            for root in image_roots
        )
        if not ok:
            raise ValueError(
                f"Imagen «{raw_path}» fuera de las raíces permitidas "
                "(vault inbound / OUTPUT). Reenvía la imagen por el chat."
            )
    return candidate


def _image_pixel_size(path: Path) -> tuple[int, int] | None:
    """Ancho×alto en píxeles; None si no se puede leer."""
    try:
        from PIL import Image

        with Image.open(path) as im:
            w, h = im.size
            if w > 0 and h > 0:
                return int(w), int(h)
    except Exception:
        pass
    try:
        raw = path.read_bytes()[:32]
        if raw[:8] == b"\x89PNG\r\n\x1a\n" and len(raw) >= 24:
            import struct

            w, h = struct.unpack(">II", raw[16:24])
            if w > 0 and h > 0:
                return int(w), int(h)
    except Exception:
        pass
    return None


def fit_inline_image_inches(
    path: Path,
    width_in: float,
    *,
    max_width_in: float = _MAX_IMAGE_WIDTH_IN,
    max_height_in: float = _MAX_IMAGE_HEIGHT_IN,
) -> tuple[float, float | None]:
    """Ajusta width/height para que la imagen quepa en una página Word.

    Returns (width_in, height_in|None). height es None si no hay aspect ratio.
    """
    width = min(max(float(width_in), 0.5), float(max_width_in))
    size = _image_pixel_size(path)
    if size is None:
        return width, None
    px_w, px_h = size
    height = width * (px_h / px_w)
    if height > max_height_in:
        scale = max_height_in / height
        width *= scale
        height = max_height_in
    if width > max_width_in:
        scale = max_width_in / width
        width *= scale
        height *= scale
    return round(width, 4), round(height, 4)


def _apply_image_sections(
    doc: Any,
    context: dict[str, Any],
    state: dict[str, Any],
    image_roots: list[Path],
) -> None:
    specs = image_render_specs(state)
    if not specs:
        return
    from docxtpl import InlineImage
    from docx.shared import Inches

    for spec in specs:
        resolved = _resolve_image_path(str(spec["path"]), image_roots)
        width_in, height_in = fit_inline_image_inches(resolved, float(spec["width_in"]))
        if height_in is not None:
            inline = InlineImage(
                doc, str(resolved), width=Inches(width_in), height=Inches(height_in)
            )
        else:
            inline = InlineImage(doc, str(resolved), width=Inches(width_in))
        assign_context_value(context, str(spec["key"]), inline)


def _safe_docx_filename(*, title: str, instance_id: str) -> str:
    """Nombre humano y estable: título normalizado + instance_id para evitar choques."""
    raw = (title or "").strip()
    ascii_title = (
        unicodedata.normalize("NFKD", raw).encode("ascii", "ignore").decode("ascii")
    )
    slug = re.sub(r"[^a-zA-Z0-9._-]+", "_", ascii_title).strip("._-")
    slug = re.sub(r"_+", "_", slug)[:72].strip("._-")
    iid = re.sub(r"[^a-zA-Z0-9_-]+", "", (instance_id or "").strip())
    if slug and iid:
        return f"{slug}_{iid}.docx"
    return f"{iid or 'documento'}.docx"


def render_instance_docx_from_uri(
    *,
    template_uri: str,
    state_json: str,
    output_root: Path,
    instance_id: str,
    title: str = "",
    period_key: str = "",
    allowed_roots: list[Path] | None = None,
    image_roots: list[Path] | None = None,
) -> dict[str, Any]:
    if allowed_roots is not None:
        template_path = assert_template_uri_readable(template_uri, allowed_roots)
    else:
        template_path = Path(template_uri).expanduser().resolve()
        if not template_path.is_file():
            raise ValueError(f"Plantilla no accesible: {template_uri}")

    state = json.loads(state_json or "{}")
    if not isinstance(state, dict):
        raise ValueError("state_json inválido")

    context = build_render_context(state)
    # setdefault no basta: la sección «titulo» vacía ya pone clave "" y tapa el title.
    title_clean = (title or "").strip()
    if not str(context.get("titulo") or "").strip() and title_clean:
        context["titulo"] = title_clean
    if not str(context.get("title") or "").strip() and title_clean:
        context["title"] = title_clean
    context.setdefault("subtitle", period_key)
    context.setdefault("period_key", period_key)
    context.setdefault("date", period_key)

    try:
        from docxtpl import DocxTemplate
    except ImportError as exc:
        raise ValueError("docxtpl no instalado (uv sync o duckops up)") from exc

    output_root.mkdir(parents=True, exist_ok=True)
    target = output_root / _safe_docx_filename(title=title, instance_id=instance_id)
    with tempfile.TemporaryDirectory(prefix="duckclaw_report_tpl_") as tmp:
        staged = Path(tmp) / f"{instance_id}_tpl.docx"
        shutil.copy2(template_path, staged)
        doc = DocxTemplate(str(staged))
        _apply_image_sections(doc, context, state, image_roots or [])
        doc.render(context, jinja_env=_lenient_jinja_env())
        doc.save(str(target))

    from duckclaw.report_engine.render_validate import find_unresolved_placeholders

    unresolved = find_unresolved_placeholders(target)
    images_embedded = 0
    try:
        import zipfile

        with zipfile.ZipFile(target) as zf:
            images_embedded = sum(
                1 for name in zf.namelist() if name.startswith("word/media/")
            )
    except Exception:
        images_embedded = len(image_render_specs(state))

    return {
        "path": str(target),
        "relative_path": target.name,
        "byte_size": target.stat().st_size,
        "format": "docx",
        "unresolved_placeholders": unresolved,
        "images_embedded": images_embedded,
    }
